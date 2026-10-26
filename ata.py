"""Gera a ata (transcrição completa) de um vídeo de reunião.

Uso:
    uv run ata.py                         # abre seletor de arquivo e pede o título
    uv run ata.py video.mp4 --titulo "Reunião X"
"""

import argparse
import os
import re
import sys
import time
from datetime import datetime
from pathlib import Path

PASTA_SAIDA = Path(__file__).parent / "atas"
MODELO = "large-v3-turbo"


def registrar_dlls_cuda():
    """Expõe as DLLs do cuBLAS/cuDNN instaladas via pip para o ctranslate2."""
    try:
        import nvidia
    except ImportError:
        return
    for base in nvidia.__path__:
        for sub in ("cublas", "cudnn"):
            bin_dir = Path(base) / sub / "bin"
            if bin_dir.is_dir():
                os.add_dll_directory(str(bin_dir))
                os.environ["PATH"] = str(bin_dir) + os.pathsep + os.environ["PATH"]


def escolher_video():
    import tkinter as tk
    from tkinter import filedialog

    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)
    caminho = filedialog.askopenfilename(
        title="Escolha o vídeo da reunião",
        filetypes=[
            ("Vídeo/Áudio", "*.mp4 *.mkv *.mov *.avi *.webm *.m4a *.mp3 *.wav *.ogg"),
            ("Todos", "*.*"),
        ],
    )
    root.destroy()
    return caminho


def carregar_modelo():
    from faster_whisper import WhisperModel

    try:
        modelo = WhisperModel(MODELO, device="cuda", compute_type="float16")
        print("Usando GPU (CUDA).")
        return modelo
    except Exception as e:
        print(f"GPU indisponível ({e}). Usando CPU, vai demorar mais.")
        return WhisperModel(MODELO, device="cpu", compute_type="int8")


def fmt_tempo(segundos):
    s = int(segundos)
    return f"{s // 3600:02d}:{s % 3600 // 60:02d}:{s % 60:02d}"


def transcrever(video):
    modelo = carregar_modelo()
    segmentos, info = modelo.transcribe(
        str(video),
        language="pt",
        vad_filter=True,
        beam_size=5,
    )
    duracao = info.duration
    print(f"Duração do vídeo: {fmt_tempo(duracao)}")

    resultado = []
    for seg in segmentos:
        texto = seg.text.strip()
        if not texto:
            continue
        resultado.append((seg.start, texto))
        pct = min(seg.end / duracao * 100, 100) if duracao else 0
        print(f"\r  {pct:5.1f}%  [{fmt_tempo(seg.end)}]", end="", flush=True)
    print()
    return resultado, duracao


def nome_arquivo(titulo):
    limpo = re.sub(r'[<>:"/\\|?*]', "", titulo).strip() or "ata"
    return f"{datetime.now():%Y-%m-%d} - {limpo}"


def salvar_txt(caminho, titulo, video, duracao, segmentos):
    linhas = [
        f"ATA - {titulo}",
        f"Data de geração: {datetime.now():%d/%m/%Y %H:%M}",
        f"Arquivo: {video.name}",
        f"Duração: {fmt_tempo(duracao)}",
        "",
        "TRANSCRIÇÃO",
        "",
    ]
    linhas += [f"[{fmt_tempo(inicio)}] {texto}" for inicio, texto in segmentos]
    caminho.write_text("\n".join(linhas) + "\n", encoding="utf-8")


def salvar_docx(caminho, titulo, video, duracao, segmentos):
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    doc.add_heading(f"Ata - {titulo}", level=0)
    info = doc.add_paragraph()
    info.add_run("Data de geração: ").bold = True
    info.add_run(f"{datetime.now():%d/%m/%Y %H:%M}\n")
    info.add_run("Arquivo: ").bold = True
    info.add_run(f"{video.name}\n")
    info.add_run("Duração: ").bold = True
    info.add_run(fmt_tempo(duracao))

    doc.add_heading("Transcrição", level=1)
    for inicio, texto in segmentos:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        marca = p.add_run(f"[{fmt_tempo(inicio)}] ")
        marca.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
        marca.font.size = Pt(9)
        p.add_run(texto)
    doc.save(caminho)


def main():
    parser = argparse.ArgumentParser(description="Gera a ata (transcrição) de um vídeo.")
    parser.add_argument("video", nargs="?", help="caminho do vídeo")
    parser.add_argument("--titulo", help="título da ata")
    args = parser.parse_args()

    video = args.video or escolher_video()
    if not video:
        print("Nenhum vídeo escolhido.")
        return 1
    video = Path(video)
    if not video.is_file():
        print(f"Arquivo não encontrado: {video}")
        return 1

    titulo = args.titulo or input("Título da ata: ").strip() or video.stem
    print(f"\nVídeo: {video}\nTítulo: {titulo}\n")

    registrar_dlls_cuda()
    inicio = time.time()
    segmentos, duracao = transcrever(video)
    if not segmentos:
        print("Nenhuma fala detectada no vídeo.")
        return 1

    PASTA_SAIDA.mkdir(exist_ok=True)
    base = PASTA_SAIDA / nome_arquivo(titulo)
    salvar_docx(base.with_suffix(".docx"), titulo, video, duracao, segmentos)
    salvar_txt(base.with_suffix(".txt"), titulo, video, duracao, segmentos)

    print(f"\nPronto em {fmt_tempo(time.time() - inicio)}.")
    print(f"  {base.with_suffix('.docx')}")
    print(f"  {base.with_suffix('.txt')}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
